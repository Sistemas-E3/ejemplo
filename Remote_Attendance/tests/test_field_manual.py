from datetime import date

from freezegun import freeze_time

from odoo.exceptions import UserError
from odoo.tests import Form, TransactionCase, new_test_user, tagged

TZ = 'Europe/Brussels'
DAY = date(2026, 9, 28)


@tagged('post_install', '-at_install')
@freeze_time('2026-09-28 18:00:00')
class TestFieldManualLoad(TransactionCase):
    """Operations types who came and their hours by hand."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.ops_user = new_test_user(cls.env, login='ops_manual', groups='Remote_Attendance.group_field_operations')
        cls.project = cls.env['project.project'].create({'name': 'Obra manual'})
        cls.other_project = cls.env['project.project'].create({'name': 'Otra obra manual'})
        Employee = cls.env['hr.employee']
        cls.supervisor = Employee.create({'name': 'Sup Manual', 'field_role': 'supervisor', 'tz': TZ})
        cls.worker_a = Employee.create({'name': 'Ana Manual', 'field_role': 'worker', 'tz': TZ,
                                        'field_supervisor_id': cls.supervisor.id})
        cls.worker_b = Employee.create({'name': 'Beto Manual', 'field_role': 'worker', 'tz': TZ,
                                        'field_supervisor_id': cls.supervisor.id})
        cls.loose = Employee.create({'name': 'Carla Manual', 'field_role': 'worker', 'tz': TZ})
        cls.supervisor.field_crew_ids = cls.worker_a | cls.worker_b

    def _form(self):
        form = Form(self.env['hr.field.manual.load'].with_user(self.ops_user))
        form.date = DAY
        form.project_id = self.project
        return form

    def _attendance(self, employee):
        return self.env['hr.attendance'].search([('employee_id', '=', employee.id), ('field_date', '=', DAY)])

    def test_supervisor_brings_crew_and_hours_are_counted(self):
        form = self._form()
        form.supervisor_id = self.supervisor
        self.assertEqual(len(form.line_ids), 3)
        with form.line_ids.edit(1) as line:       # Ana: 7 to 19, two hours of overtime
            self.assertEqual(line.hours, 9.0)
            line.check_out = 19.0
            self.assertEqual(line.hours, 9.0)
            self.assertEqual(line.overtime_hours, 2.0)
        with form.line_ids.edit(2) as line:       # Beto: types 5 hours, exit moves to 12:00
            line.hours = 5.0
            self.assertEqual(line.check_out, 13.0)  # 5 hours plus the lunch
            line.lunch = False
            line.hours = 5.0
            self.assertEqual(line.check_out, 12.0)
        with form.line_ids.edit(0) as line:       # the supervisor did not come
            line.present = False
        wizard = form.save()
        self.assertEqual(wizard.present_count, 2)
        wizard.action_confirm()

        self.assertFalse(self._attendance(self.supervisor))
        ana = self._attendance(self.worker_a)
        self.assertEqual(ana.field_regular_hours, 9.0)
        self.assertEqual(ana.field_overtime_hours, 2.0)
        self.assertEqual(ana.field_overtime_state, 'approved')
        self.assertEqual(ana.field_loaded_by_id, self.ops_user)
        self.assertEqual(ana.field_supervisor_id, self.supervisor)
        self.assertEqual(sum(ana.field_timesheet_ids.mapped('unit_amount')), 11.0)
        self.assertEqual(ana.field_timesheet_ids.project_id, self.project)
        beto = self._attendance(self.worker_b)
        self.assertEqual(beto.field_regular_hours, 5.0)
        self.assertFalse(beto.field_overtime_state)
        self.assertEqual(beto.field_timesheet_ids.unit_amount, 5.0)

    def test_no_lunch_full_day_and_other_project(self):
        form = self._form()
        with form.line_ids.new() as line:
            line.employee_id = self.loose
            self.assertFalse(line.supervisor_id, "Carla has no crew")
            line.supervisor_id = self.supervisor
            line.lunch = False
            line.project_id = self.other_project
        form.save().action_confirm()
        attendance = self._attendance(self.loose)
        self.assertTrue(attendance.field_skip_lunch)
        self.assertEqual(attendance.field_regular_hours, 10.0)
        self.assertEqual(attendance.field_project_id, self.other_project)
        self.assertEqual(attendance.field_supervisor_id, self.supervisor)

    def test_all_field_staff_and_existing_attendance_is_skipped(self):
        form = self._form()
        wizard = form.save()
        wizard.action_add_field_staff()
        self.assertTrue({self.supervisor, self.worker_a, self.worker_b, self.loose} <= set(wizard.line_ids.employee_id))
        wizard.line_ids.filtered(lambda l: l.employee_id not in (self.worker_a | self.loose)).present = False
        with self.assertRaises(UserError):  # Carla has no crew: who brings her must be chosen
            wizard.action_confirm()
        wizard.line_ids.filtered(lambda l: l.employee_id == self.loose).supervisor_id = self.supervisor
        self.assertEqual(wizard.line_ids.filtered(lambda l: l.employee_id == self.worker_a).supervisor_id, self.supervisor)
        wizard.action_confirm()
        self.assertTrue(self._attendance(self.worker_a))

        again = self._form().save()
        again.line_ids = [(0, 0, {'employee_id': self.worker_a.id}), (0, 0, {'employee_id': self.worker_b.id})]
        result = again.action_confirm()
        self.assertIn('Ana Manual', result['params']['message'])
        self.assertEqual(len(self._attendance(self.worker_a)), 1)
        self.assertTrue(self._attendance(self.worker_b))

    def test_errors(self):
        wizard = self._form().save()
        with self.assertRaises(UserError):
            wizard.action_confirm()
        wizard.line_ids = [(0, 0, {'employee_id': self.worker_a.id, 'check_in': 17.0, 'check_out': 9.0})]
        with self.assertRaises(UserError):
            wizard.action_confirm()
        wizard.line_ids = [(5,), (0, 0, {'employee_id': self.worker_a.id}), (0, 0, {'employee_id': self.worker_a.id})]
        with self.assertRaises(UserError):
            wizard.action_confirm()
        wizard.write({'date': date(2026, 9, 30), 'line_ids': [(5,), (0, 0, {'employee_id': self.worker_b.id})]})
        with self.assertRaises(UserError):
            wizard.action_confirm()


@tagged('post_install', '-at_install')
@freeze_time('2026-09-28 18:00:00')
class TestFieldTimesheetList(TestFieldManualLoad):
    """Saved lists of hours exported to Excel with the timesheet import columns."""

    def test_list_and_excel(self):
        import base64
        import io
        import openpyxl
        from ..models.hr_field_timesheet_list import IMPORT_HEADERS

        form = self._form()
        form.supervisor_id = self.supervisor
        with form.line_ids.edit(1) as line:
            line.check_out = 19.0          # Ana: 9 + 2 approved overtime
        with form.line_ids.edit(0) as line:
            line.present = False
        form.save().action_confirm()
        self._attendance(self.worker_b).field_state = 'review'   # not accepted: stays out of the list

        Lists = self.env['hr.field.timesheet.list'].with_user(self.ops_user)
        record = Lists.create({'date_from': DAY, 'date_to': DAY, 'supervisor_id': self.supervisor.id})
        self.assertIn('Sup Manual', record.name)
        record.action_generate()
        self.assertEqual(record.line_ids.employee_id, self.worker_a)
        self.assertEqual(record.line_ids.hours, 11.0)
        self.assertEqual(record.line_ids.overtime_hours, 2.0)
        self.assertEqual(record.line_ids.project_id, self.project)

        action = record.action_export_excel()
        self.assertEqual(record.state, 'exported')
        self.assertIn('/excel_file/', action['url'])
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(record.excel_file)))
        sheet = book['Hojas de horas']
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), IMPORT_HEADERS)
        self.assertEqual(rows[1][1:], ('Ana Manual', 'Obra manual', None, 'Asistencia en obra', 11))
        self.assertEqual(rows[1][0].date(), DAY)
        detail = list(book['Detalle'].iter_rows(values_only=True))
        self.assertEqual(detail[1][3], 'Sup Manual')
        self.assertEqual(detail[1][5:7], ('07:00', '19:00'))

        other = Lists.create({'date_from': DAY, 'date_to': DAY, 'project_id': self.other_project.id})
        with self.assertRaises(UserError):
            other.action_export_excel()

    def test_old_lists_are_deleted(self):
        Lists = self.env['hr.field.timesheet.list']
        old, recent = Lists.create([{'date_from': DAY, 'date_to': DAY}, {'date_from': DAY, 'date_to': DAY}])
        self.env.cr.execute("UPDATE hr_field_timesheet_list SET create_date = %s WHERE id = %s",
                            ('2025-06-01 00:00:00', old.id))
        old.invalidate_recordset(['create_date'])
        Lists._cron_delete_old()
        self.assertFalse(old.exists())
        self.assertTrue(recent.exists())
        self.env['ir.config_parameter'].sudo().set_param('Remote_Attendance.list_keep_days', '0')
        self.env.cr.execute("UPDATE hr_field_timesheet_list SET create_date = %s WHERE id = %s",
                            ('2024-01-01 00:00:00', recent.id))
        recent.invalidate_recordset(['create_date'])
        Lists._cron_delete_old()
        self.assertTrue(recent.exists())
